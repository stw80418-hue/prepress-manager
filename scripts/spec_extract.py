"""Extract printer text plus candidates. Human/Codex resolves language, units, scope.

Never silently promotes ambiguous OCR/regex output into authoritative rules.
"""
import argparse
import json
import re
from pathlib import Path


def extract(path):
    ext = path.suffix.lower()
    notes = []
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as pdf:
            chunks = [(f"page {i+1}", p.get_text()) for i,p in enumerate(pdf)]
        if any(not t.strip() for _,t in chunks):
            notes.append("部分頁面沒有文字層，必須閱讀原頁面／OCR 並校對數值。")
    elif ext == ".docx":
        from docx import Document
        d = Document(path)
        chunks = [(f"paragraph {i+1}", p.text) for i,p in enumerate(d.paragraphs)]
        for i,table in enumerate(d.tables):
            chunks.extend((f"table {i+1} row {j+1}", " | ".join(c.text for c in row.cells)) for j,row in enumerate(table.rows))
        notes.append("DOCX 圖片、文字方塊、頁首／頁尾需目視核對。")
    elif ext == ".txt":
        chunks = [(f"line {i+1}",line) for i,line in enumerate(path.read_text(encoding="utf-8-sig").splitlines())]
    elif ext in [".png",".jpg",".jpeg",".tif",".tiff",".webp"]:
        chunks = []
        notes.append("圖片規格須由 Codex 讀圖或本機 OCR 擷取，再對照原圖逐項驗證數值與單位；此腳本不捏造 OCR 結果。")
    else:
        raise ValueError("Supported specs: PDF, DOCX, UTF-8 TXT and images")
    patterns = {
        "bleed_mm":r"(?:出血|bleed)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*(?:mm|毫米|公釐)",
        "safe_margin_mm":r"(?:安全距離|安全邊界|safe\s*margin)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*(?:mm|毫米|公釐)",
        "minimum_font_size_pt":r"(?:最小字級|最小文字|minimum\s*(?:font|text)(?:\s*size)?)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*(?:pt|點)",
        "minimum_line_width_pt":r"(?:最小線寬|線寬|minimum\s*line\s*width)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*(?:pt|點)",
        "maximum_tac_percent":r"(?:TAC|總墨量)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*[%％]",
        "minimum_image_ppi":r"(?:解析度|resolution|最低解析度)[^\n\d]{0,20}(\d+(?:\.\d+)?)\s*(?:ppi|dpi)"
    }
    candidates=[]
    for location,txt in chunks:
        for key,pattern in patterns.items():
            for match in re.finditer(pattern,txt,re.I):
                candidates.append({"key":key,"value":float(match[1]),"quote":match[0],"location":location})
    conflicts=[]
    for key in patterns:
        vals={c["value"] for c in candidates if c["key"]==key}
        if len(vals)>1:
            conflicts.append(f"{key}: {sorted(vals)}；可能依紙材／品項分流，必須釐清。")
    spec = {"schema_version":1,"source":str(path.resolve()),"review_status":"needs_review","overrides":{},"evidence":{},
            "manual_requirements":[],"candidates":candidates,"conflicts":conflicts,
            "notes":notes+["先閱讀完整文字與原始排版；負向條件、雙面、不同紙材與單位不可只靠 regex。"]}
    return spec, "\n\n".join(f"[{loc}]\n{txt}" for loc,txt in chunks)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("input",type=Path)
    ap.add_argument("--out",required=True,type=Path,help="A NEW output directory")
    args=ap.parse_args()
    spec,text=extract(args.input)
    args.out.mkdir(parents=True,exist_ok=False)
    (args.out/"PRINT_SPEC.json").write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding="utf-8")
    (args.out/"PRINT_SPEC_SOURCE.txt").write_text(text,encoding="utf-8")
    print(args.out/"PRINT_SPEC.json")


if __name__=="__main__":
    main()
